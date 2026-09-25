# Génère le classeur de relecture native somali / afar à partir de translations-review/data.json
# (produit par scripts/translations_review_export.mjs). Sortie : translations-review/relecture-so-aa-<date>.xlsx
import json, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

data = json.load(open('translations-review/data.json', encoding='utf-8'))
today = datetime.date.today().isoformat()
out = f'translations-review/relecture-so-aa-{today}.xlsx'

FONT = 'Arial'
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_title = Font(name=FONT, size=14, bold=True, color='2D6410')
fill_head = PatternFill('solid', fgColor='526500')
fill_fix = PatternFill('solid', fgColor='FFFF00')      # colonnes à remplir par le relecteur
fill_bad = PatternFill('solid', fgColor='FDE2E2')      # afar = copie du somali (à refaire)
fill_miss = PatternFill('solid', fgColor='FFF3CD')     # traduction absente
thin = Side(style='thin', color='D2E095')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# ── Lisez-moi ──
ws = wb.active; ws.title = 'Lisez-moi'
ws['A1'] = 'Hornafresh — relecture native des traductions somali (so) et afar (aa)'; ws['A1'].font = f_title
lines = [
    f'Généré le {today} depuis la base de production (table ui_translations + product_translations).',
    '',
    'COMMENT RELIRE',
    '1. Onglet « Interface » : une ligne par texte. Colonne C = français de référence, D = anglais (aide).',
    '2. Colonne E = somali actuel. Si le somali est correct, ne rien écrire. Sinon, écrire la bonne version dans F « Somali corrigé » (jaune).',
    '3. Colonne G = afar actuel. ATTENTION : les cellules en rouge clair sont une COPIE DU SOMALI (ce n\'est pas de l\'afar) — à traduire en afar dans H « Afar corrigé » (jaune).',
    '   Beaucoup d\'autres cellules afar sont du somali abrégé ou de l\'oromo : les traduire aussi dans H.',
    '4. Les cellules orange pâle = traduction absente : à écrire dans la colonne corrigée.',
    '5. Ne pas modifier les colonnes A à D. Ne pas supprimer ni déplacer de lignes.',
    '6. Conserver les éléments techniques tels quels : émojis, « {count} », chiffres, « Fdj », noms « Hornafresh » / « Waafi ».',
    '7. Conventions : les exemples de champs commencent par « Tusaale: » en somali ; choisir un équivalent afar et l\'utiliser partout.',
    '8. Onglet « Produits » : noms et descriptions des produits, même principe (colonnes jaunes à remplir).',
    '',
    'IMPORT DES CORRECTIONS (côté Hornafresh)',
    'node scripts/translations_review_import.mjs translations-review/relecture-so-aa-' + today + '.xlsx --dry   (aperçu)',
    'node scripts/translations_review_import.mjs translations-review/relecture-so-aa-' + today + '.xlsx         (application)',
    '',
    'SUIVI',
]
for i, l in enumerate(lines, start=3):
    ws.cell(row=i, column=1, value=l).font = f_bold if l.isupper() and l else f_norm
r = 3 + len(lines)
n_ui = len(data['ui'])
ws.cell(row=r, column=1, value='Textes d\'interface').font = f_norm; ws.cell(row=r, column=2, value=n_ui).font = f_norm
ws.cell(row=r+1, column=1, value='Somali absent').font = f_norm; ws.cell(row=r+1, column=2, value=sum(1 for x in data['ui'] if not x['so'])).font = f_norm
ws.cell(row=r+2, column=1, value='Afar absent').font = f_norm; ws.cell(row=r+2, column=2, value=sum(1 for x in data['ui'] if not x['aa'])).font = f_norm
ws.cell(row=r+3, column=1, value='Afar = copie du somali (à refaire)').font = f_norm; ws.cell(row=r+3, column=2, value=sum(1 for x in data['ui'] if x['aa_copy_of_so'])).font = f_norm
ws.cell(row=r+4, column=1, value='Corrections somali saisies').font = f_bold; ws.cell(row=r+4, column=2, value=f'=COUNTA(Interface!F2:F{n_ui+1})').font = f_bold
ws.cell(row=r+5, column=1, value='Corrections afar saisies').font = f_bold; ws.cell(row=r+5, column=2, value=f'=COUNTA(Interface!H2:H{n_ui+1})').font = f_bold
ws.cell(row=r+6, column=1, value='Corrections produits saisies').font = f_bold; ws.cell(row=r+6, column=2, value="=COUNTA(Produits!F2:F200)+COUNTA(Produits!H2:H200)+COUNTA(Produits!J2:J200)+COUNTA(Produits!L2:L200)").font = f_bold
ws.column_dimensions['A'].width = 120; ws.column_dimensions['B'].width = 14

# ── Interface ──
wi = wb.create_sheet('Interface')
heads = ['Clé', 'Fichier (contexte)', 'Français (référence)', 'Anglais', 'Somali actuel', 'Somali corrigé', 'Afar actuel', 'Afar corrigé', 'Remarque']
for c, h in enumerate(heads, start=1):
    cell = wi.cell(row=1, column=c, value=h); cell.font = f_head; cell.fill = fill_head; cell.alignment = Alignment(vertical='center', wrap_text=True); cell.border = border
widths = [34, 30, 48, 40, 40, 40, 40, 40, 30]
for i, w in enumerate(widths, start=1): wi.column_dimensions[get_column_letter(i)].width = w
for rIdx, row in enumerate(data['ui'], start=2):
    vals = [row['key'], row['file'], row['fr'], row['en'], row['so'], '', row['aa'], '', '']
    note = []
    if row['aa_copy_of_so']: note.append('Afar = copie du somali')
    if not row['so']: note.append('somali absent')
    if not row['aa']: note.append('afar absent')
    if not row['fr']: note.append('référence FR à retrouver dans le code')
    vals[8] = ' · '.join(note)
    for c, v in enumerate(vals, start=1):
        cell = wi.cell(row=rIdx, column=c, value=v); cell.font = f_norm; cell.alignment = wrap; cell.border = border
    wi.cell(row=rIdx, column=6).fill = fill_fix; wi.cell(row=rIdx, column=8).fill = fill_fix
    if row['aa_copy_of_so']: wi.cell(row=rIdx, column=7).fill = fill_bad
    if not row['so']: wi.cell(row=rIdx, column=5).fill = fill_miss
    if not row['aa'] and not row['aa_copy_of_so']: wi.cell(row=rIdx, column=7).fill = fill_miss
wi.freeze_panes = 'C2'; wi.auto_filter.ref = f'A1:I{n_ui+1}'

# ── Produits ──
wp = wb.create_sheet('Produits')
ph = ['Produit (id)', 'Nom FR', 'Description FR', 'Nom EN', 'Nom SO actuel', 'Nom SO corrigé', 'Nom AA actuel', 'Nom AA corrigé', 'Description SO actuelle', 'Description SO corrigée', 'Description AA actuelle', 'Description AA corrigée']
for c, h in enumerate(ph, start=1):
    cell = wp.cell(row=1, column=c, value=h); cell.font = f_head; cell.fill = fill_head; cell.alignment = Alignment(vertical='center', wrap_text=True); cell.border = border
pw = [12, 26, 40, 26, 26, 26, 26, 26, 40, 40, 40, 40]
for i, w in enumerate(pw, start=1): wp.column_dimensions[get_column_letter(i)].width = w
for rIdx, p in enumerate(data['products'], start=2):
    vals = [p['product_id'], p['fr_name'], p['fr_description'], p.get('en_name', ''), p.get('so_name', ''), '', p.get('aa_name', ''), '', p.get('so_description', ''), '', p.get('aa_description', ''), '']
    for c, v in enumerate(vals, start=1):
        cell = wp.cell(row=rIdx, column=c, value=v); cell.font = f_norm; cell.alignment = wrap; cell.border = border
    for c in (6, 8, 10, 12): wp.cell(row=rIdx, column=c).fill = fill_fix
    if p.get('aa_name') and p.get('aa_name') == p.get('so_name'): wp.cell(row=rIdx, column=7).fill = fill_bad
    if p.get('aa_description') and p.get('aa_description') == p.get('so_description'): wp.cell(row=rIdx, column=11).fill = fill_bad
wp.freeze_panes = 'C2'

wb.save(out)
print(out)
